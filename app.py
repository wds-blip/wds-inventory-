import csv, io, os, secrets, time
from functools import wraps
from decimal import Decimal, InvalidOperation
from flask import Flask, jsonify, request, session, render_template_string, redirect
from werkzeug.exceptions import HTTPException
import psycopg
from psycopg.rows import dict_row
from openpyxl import load_workbook

app=Flask(__name__)
app.secret_key=os.environ.get('SESSION_SECRET','local-development-only-change-me')
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=os.getenv('COOKIE_SECURE','true').lower()=='true', PERMANENT_SESSION_LIFETIME=60*60*10)
DB=os.environ.get('DATABASE_URL','').replace('postgres://','postgresql://',1)
app.config['MAX_CONTENT_LENGTH']=10*1024*1024
ADMIN_EMAIL=os.environ.get('ADMIN_EMAIL','wds@telus.net').strip().lower()
ADMIN_PASSWORD=os.environ.get('ADMIN_PASSWORD','')
LOGIN_ATTEMPTS={}

def db():
    if not DB: raise RuntimeError('DATABASE_URL is required')
    return psycopg.connect(DB, row_factory=dict_row, connect_timeout=10)

def init_db():
    with db() as conn:
        conn.execute('''CREATE TABLE IF NOT EXISTS products (
          id BIGSERIAL PRIMARY KEY, brand TEXT NOT NULL DEFAULT '', category TEXT NOT NULL DEFAULT '', part_number TEXT NOT NULL DEFAULT '',
          name TEXT NOT NULL, purchase_price NUMERIC(12,2), list_price NUMERIC(12,2), quantity INTEGER NOT NULL DEFAULT 0,
          location TEXT NOT NULL DEFAULT '', memo TEXT NOT NULL DEFAULT '', updated_at TIMESTAMPTZ NOT NULL DEFAULT now())''')
        conn.execute('''CREATE TABLE IF NOT EXISTS inventory_movements (
          id BIGSERIAL PRIMARY KEY, product_id BIGINT NOT NULL REFERENCES products(id), delta INTEGER NOT NULL,
          quantity_after INTEGER NOT NULL, note TEXT NOT NULL DEFAULT '', changed_by TEXT NOT NULL, created_at TIMESTAMPTZ NOT NULL DEFAULT now())''')

def require_login(fn):
    @wraps(fn)
    def wrapped(*args,**kwargs):
        if not session.get('admin'): return jsonify({'error':'Login required'}),401
        return fn(*args,**kwargs)
    return wrapped

def csrf_ok():
    return secrets.compare_digest(session.get('csrf',''),request.headers.get('X-CSRF-Token',''))

def clean_price(value):
    if value in (None,''): return None
    try:
        n=Decimal(str(value))
        if n<0 or n>Decimal('9999999999'): raise ValueError()
        return n.quantize(Decimal('.01'))
    except (InvalidOperation,ValueError): raise ValueError('Enter a valid non-negative price.')

@app.get('/health')
def health(): return {'ok':True}

@app.get('/login')
def login_page():
    if session.get('admin'): return redirect('/')
    return render_template_string(LOGIN)

@app.post('/login')
def login():
    ip=request.headers.get('X-Forwarded-For',request.remote_addr or 'unknown').split(',')[0]
    now=time.time(); attempts=[x for x in LOGIN_ATTEMPTS.get(ip,[]) if now-x<900]
    if len(attempts)>=10: return render_template_string(LOGIN,error='Too many attempts. Wait 15 minutes and try again.'),429
    email=request.form.get('email','').strip().lower(); password=request.form.get('password','')
    if not ADMIN_PASSWORD or not secrets.compare_digest(email,ADMIN_EMAIL) or not secrets.compare_digest(password,ADMIN_PASSWORD):
        attempts.append(now); LOGIN_ATTEMPTS[ip]=attempts
        return render_template_string(LOGIN,error='Email or password is incorrect.'),401
    LOGIN_ATTEMPTS.pop(ip,None); session.clear(); session.permanent=True; session['admin']=ADMIN_EMAIL; session['csrf']=secrets.token_urlsafe(32)
    return redirect('/')

@app.post('/logout')
@require_login
def logout():
    supplied=request.headers.get('X-CSRF-Token',request.form.get('csrf',''))
    if not secrets.compare_digest(session.get('csrf',''),supplied): return jsonify({'error':'Invalid request token'}),403
    session.clear(); return redirect('/login')

@app.get('/')
def home():
    if not session.get('admin'): return redirect('/login')
    return render_template_string(APP,csrf=session['csrf'],email=session['admin'])

@app.get('/api/products')
@require_login
def products():
    q=request.args.get('q','').strip()[:150]
    params=[]; where=''
    if q:
        where='WHERE brand ILIKE %s OR category ILIKE %s OR part_number ILIKE %s OR name ILIKE %s OR location ILIKE %s'
        pattern='%'+q+'%'; params=[pattern]*5
    with db() as conn:
        rows=conn.execute(f'''SELECT id,brand,category,part_number,name,purchase_price,list_price,quantity,location,memo
          FROM products {where} ORDER BY brand,name LIMIT 500''',params).fetchall()
        totals=conn.execute('''SELECT count(*) AS items, coalesce(sum(quantity),0) AS units,
          coalesce(sum(quantity*coalesce(purchase_price,0)),0) AS cost_value,
          coalesce(sum(quantity*coalesce(list_price,0)),0) AS list_value,
          count(*) FILTER (WHERE purchase_price IS NULL OR list_price IS NULL) AS missing_prices FROM products''').fetchone()
    return jsonify({'products':rows,'totals':totals})

@app.post('/api/import-workbook')
@require_login
def import_workbook():
    if not csrf_ok(): return jsonify({'error':'Invalid request token'}),403
    upload=request.files.get('workbook')
    if not upload or not upload.filename.lower().endswith('.xlsx'):
        return jsonify({'error':'Choose the WDS inventory .xlsx workbook.'}),400
    try:
        workbook=load_workbook(upload.stream,data_only=True,read_only=True)
        inventory=workbook['Inventory Count']; audit=workbook['Source Audit']; sales={}
        for row in audit.iter_rows(min_row=2,values_only=True):
            if row[1]: sales[row[1]]=row[4]
        imported=[]
        for row in inventory.iter_rows(min_row=2,values_only=True):
            if not row[3]: continue
            imported.append((row[0] or '',row[1] or '',str(row[2] or ''),str(row[3]),row[5],sales.get(row[16]),int(row[6] or row[9] or 0),row[4] or '',row[18] or ''))
        workbook.close()
        if not imported: return jsonify({'error':'No product rows were found in Inventory Count.'}),400
    except Exception:
        return jsonify({'error':'Could not read the workbook. Use the provided WDS inventory workbook.'}),400
    try:
        with db() as conn:
            n=conn.execute('SELECT count(*) AS n FROM products').fetchone()['n']
            if n: return jsonify({'error':'The catalog already contains products. Import is only enabled for an empty inventory.'}),409
            with conn.cursor() as cur:
                cur.executemany('''INSERT INTO products (brand,category,part_number,name,purchase_price,list_price,quantity,location,memo)
                  VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)''',imported)
        return jsonify({'ok':True,'imported':len(imported)})
    except Exception:
        app.logger.exception('Workbook import failed')
        return jsonify({'error':'The catalog import failed. No products were loaded.'}),500

@app.patch('/api/products/<int:pid>')
@require_login
def update_product(pid):
    if not csrf_ok(): return jsonify({'error':'Invalid request token'}),403
    data=request.get_json(silent=True) or {}
    allowed={'purchase_price','list_price','location'}; updates={k:data[k] for k in allowed if k in data}
    if not updates: return jsonify({'error':'No editable fields provided'}),400
    try:
        if 'purchase_price' in updates: updates['purchase_price']=clean_price(updates['purchase_price'])
        if 'list_price' in updates: updates['list_price']=clean_price(updates['list_price'])
    except ValueError as e: return jsonify({'error':str(e)}),400
    updates['updated_at']=None
    cols=', '.join(f'{k}=%s' for k in updates if k!='updated_at')
    vals=[v for k,v in updates.items() if k!='updated_at']+[pid]
    with db() as conn:
        row=conn.execute(f'UPDATE products SET {cols}, updated_at=now() WHERE id=%s RETURNING id',vals).fetchone()
    if not row: return jsonify({'error':'Product not found'}),404
    return jsonify({'ok':True})

@app.post('/api/products/<int:pid>/quantity')
@require_login
def adjust_quantity(pid):
    if not csrf_ok(): return jsonify({'error':'Invalid request token'}),403
    data=request.get_json(silent=True) or {}
    try: delta=int(data.get('delta'))
    except (TypeError,ValueError): return jsonify({'error':'Quantity change must be a whole number.'}),400
    if not -1000000<=delta<=1000000 or delta==0: return jsonify({'error':'Enter a non-zero quantity change.'}),400
    note=str(data.get('note',''))[:200]
    with db() as conn:
        row=conn.execute('UPDATE products SET quantity=quantity+%s,updated_at=now() WHERE id=%s AND quantity+%s>=0 RETURNING quantity',(delta,pid,delta)).fetchone()
        if row: conn.execute('INSERT INTO inventory_movements(product_id,delta,quantity_after,note,changed_by) VALUES (%s,%s,%s,%s,%s)',(pid,delta,row['quantity'],note,session['admin']))
    if not row: return jsonify({'error':'Product not found or quantity cannot go below zero.'}),400
    return jsonify({'ok':True,'quantity':row['quantity']})

@app.errorhandler(Exception)
def error(e):
    if isinstance(e,HTTPException): return e
    app.logger.exception('Request failed')
    return jsonify({'error':'The request could not be completed.'}),500

LOGIN='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>WDS Inventory · Sign in</title><style>
*{box-sizing:border-box}body{margin:0;background:#f1f5f9;font:16px system-ui;color:#172033;min-height:100vh;display:grid;place-items:center}.card{width:min(420px,92vw);background:white;border:1px solid #dce3ee;border-radius:18px;padding:32px;box-shadow:0 16px 50px #10213b14}.brand{font-size:13px;letter-spacing:.12em;color:#59708e;font-weight:800}.logo{width:45px;height:45px;border-radius:14px;background:#103a62;color:white;display:grid;place-items:center;font-weight:800;margin-bottom:20px}h1{font-size:27px;margin:12px 0 6px}.muted{color:#6a778a;margin:0 0 24px}label{display:block;font-size:13px;font-weight:700;margin:16px 0 7px}input{width:100%;padding:12px;border:1px solid #cad5e3;border-radius:9px;background:#fff;font:inherit}button{width:100%;border:0;border-radius:9px;padding:13px;margin-top:20px;background:#103a62;color:#fff;font-weight:750;font-size:15px;cursor:pointer}.error{background:#fff1f0;color:#a52520;border-radius:8px;padding:10px;margin:12px 0}</style><main class="card"><div class="logo">W</div><div class="brand">WESTERN DOOR SOLUTIONS</div><h1>Inventory sign in</h1><p class="muted">Sign in to manage the WDS inventory.</p>{% if error %}<div class="error">{{error}}</div>{% endif %}<form method="post"><label for="email">Email</label><input id="email" name="email" type="email" autocomplete="username" required><label for="password">Password</label><input id="password" name="password" type="password" autocomplete="current-password" required><button>Sign in</button></form></main></html>'''

APP='''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>WDS Inventory</title><script src="https://cdn.jsdelivr.net/npm/tesseract.js@5/dist/tesseract.min.js" crossorigin="anonymous"></script><style>
:root{--navy:#123b62;--ink:#172033;--muted:#69788b;--line:#e2e8f0;--bg:#f4f7fb;--green:#13815d}*{box-sizing:border-box}body{margin:0;background:var(--bg);font:14px system-ui,-apple-system,Segoe UI,sans-serif;color:var(--ink)}header{background:linear-gradient(120deg,#0d3154,#19527e);color:white;padding:22px max(22px,calc((100vw - 1400px)/2));display:flex;align-items:center;justify-content:space-between;gap:20px}.brand{display:flex;gap:12px;align-items:center}.logo{width:42px;height:42px;border-radius:13px;background:#ffffff22;display:grid;place-items:center;font-weight:900;font-size:18px}.brand h1{font-size:19px;margin:0}.brand small{color:#c8d8e8}.user{display:flex;align-items:center;gap:14px}.user button{border:1px solid #ffffff55;background:transparent;color:#fff;border-radius:8px;padding:8px 12px;cursor:pointer}.wrap{max-width:1400px;margin:26px auto;padding:0 22px}.cards{display:grid;grid-template-columns:repeat(4,1fr);gap:14px}.card{background:white;border:1px solid var(--line);border-radius:13px;padding:17px}.card label{font-size:12px;color:var(--muted);font-weight:650}.value{font-size:25px;font-weight:800;margin-top:6px;letter-spacing:-.03em}.hint{font-size:11px;color:var(--muted);margin-top:4px}.tools{margin:20px 0 12px;display:flex;gap:10px;align-items:center;flex-wrap:wrap}.search{flex:1;min-width:240px;position:relative}.search input{width:100%;padding:12px 14px 12px 42px;border-radius:10px;border:1px solid #ced8e4;background:white;font:inherit}.search span{position:absolute;left:14px;top:12px;color:#8391a3;font-size:17px}.btn{background:var(--navy);color:white;border:0;border-radius:9px;padding:11px 14px;font-weight:700;cursor:pointer}.btn.light{background:white;color:var(--navy);border:1px solid #ccd7e4}.btn:disabled{opacity:.5}.wrapsmall{font-size:12px;color:var(--muted)}.tablebox{background:white;border:1px solid var(--line);border-radius:13px;overflow:hidden}table{width:100%;border-collapse:collapse}th{text-align:left;background:#f8fafc;color:#627187;font-size:11px;letter-spacing:.04em;text-transform:uppercase;padding:12px 10px;border-bottom:1px solid var(--line)}td{padding:11px 10px;border-bottom:1px solid #edf1f5;vertical-align:middle}tr:last-child td{border:0}.part{font-family:ui-monospace,monospace;font-size:12px;font-weight:750;color:#164a74}.name{font-weight:650}.sub{color:var(--muted);font-size:11px;margin-top:3px}.qty{white-space:nowrap}.qty button{width:27px;height:27px;background:#eff4f8;border:1px solid #d7e0e9;border-radius:7px;cursor:pointer;font-weight:800}.qty strong{display:inline-block;min-width:34px;text-align:center;font-size:14px}.price{width:85px;padding:7px;border:1px solid #d8e1eb;border-radius:7px;text-align:right;background:#fff}.missing{color:#b45309;font-size:11px}.count{font-variant-numeric:tabular-nums}.progress{color:var(--muted);margin:10px 2px}.empty{text-align:center;padding:38px;color:var(--muted)}.modal{position:fixed;inset:0;background:#07182ca6;display:none;place-items:center;padding:20px;z-index:2}.modal.open{display:grid}.dialog{max-width:620px;width:100%;background:#fff;border-radius:16px;padding:20px}.dialog h2{margin:0 0 8px}.dialog p{color:var(--muted)}video{width:100%;max-height:360px;background:#0b1724;border-radius:12px}#photo{max-width:100%;display:none;margin-top:12px;border-radius:10px}.actions{display:flex;gap:9px;flex-wrap:wrap;margin-top:14px}.close{float:right;border:0;background:#eef2f6;border-radius:8px;padding:7px 11px;cursor:pointer}.note{font-size:11px;color:var(--muted);margin-top:12px}.toast{position:fixed;bottom:20px;right:20px;background:#142e49;color:#fff;padding:12px 16px;border-radius:10px;display:none;max-width:90vw;z-index:3}@media(max-width:850px){.cards{grid-template-columns:repeat(2,1fr)}.tablebox{overflow:auto}table{min-width:870px}}@media(max-width:500px){header{padding:16px}.wrap{padding:0 12px}.cards{gap:8px}.card{padding:12px}.value{font-size:20px}.user span{display:none}}
</style></head><body><header><div class="brand"><div class="logo">W</div><div><h1>WDS Inventory</h1><small>Stock, product costs and list values</small></div></div><div class="user"><span>{{email}}</span><form method="post" action="/logout"><input type="hidden" name="csrf" value="{{csrf}}"><button id="logout">Sign out</button></form></div></header><main class="wrap"><section class="cards"><div class="card"><label>Products</label><div class="value" id="items">—</div><div class="hint">Active catalog items</div></div><div class="card"><label>Total quantity</label><div class="value" id="units">—</div><div class="hint">Units currently on hand</div></div><div class="card"><label>Inventory cost value</label><div class="value" id="cost">—</div><div class="hint">Quantity × purchase price</div></div><div class="card"><label>Potential list value</label><div class="value" id="list">—</div><div class="hint">Quantity × list price</div></div></section><div class="tools"><div class="search"><span>⌕</span><input id="q" placeholder="Search brand, part number, product or location…" autocomplete="off"></div><button class="btn light" id="scan">📷 Scan part number</button><label class="btn light" for="workbook">Import workbook</label><input type="file" id="workbook" accept=".xlsx" hidden><button class="btn light" id="refresh">Refresh</button></div><div class="progress" id="status">Loading products…</div><div class="tablebox"><table><thead><tr><th>Part number</th><th>Product</th><th>Location</th><th>Quantity</th><th>Purchase price</th><th>List price</th></tr></thead><tbody id="rows"></tbody></table></div><p class="note">Inventory values include priced items only. Blank prices are flagged for completion. Quantity changes are recorded in the inventory history.</p></main><div class="modal" id="modal"><section class="dialog"><button class="close" id="close">Close</button><h2>Scan part number</h2><p>Take a photo or upload an image. Text recognition runs in your browser; detected text fills the search box.</p><video id="video" autoplay playsinline></video><canvas id="canvas" hidden></canvas><img id="photo"><div class="actions"><button class="btn" id="capture">Take photo</button><label class="btn light" for="file">Choose image</label><input type="file" id="file" accept="image/*" capture="environment" hidden><button class="btn light" id="ocr">Read part number</button></div><div class="progress" id="scanstatus"></div></section></div><div class="toast" id="toast"></div><script>
const CSRF={{csrf|tojson}};let all=[];const fmt=n=>new Intl.NumberFormat('en-CA',{style:'currency',currency:'CAD'}).format(Number(n||0));
document.querySelector('#workbook').onchange=async e=>{const file=e.target.files[0];if(!file)return;if(!confirm(`Import the catalog from ${file.name}? This only works on an empty inventory.`)){e.target.value='';return}const form=new FormData();form.append('workbook',file);document.querySelector('#status').textContent='Importing workbook…';try{const r=await fetch('/api/import-workbook',{method:'POST',headers:{'X-CSRF-Token':CSRF},body:form});const d=await r.json();if(!r.ok)throw Error(d.error||'Import failed');toast(`${d.imported.toLocaleString()} products imported.`);await load()}catch(err){toast(err.message);document.querySelector('#status').textContent='Import failed.'}e.target.value=''};
async function load(){const q=document.querySelector('#q').value;const r=await fetch('/api/products?q='+encodeURIComponent(q));if(!r.ok){document.querySelector('#status').textContent='Could not load inventory.';return}const d=await r.json();all=d.products;document.querySelector('#items').textContent=Number(d.totals.items).toLocaleString();document.querySelector('#units').textContent=Number(d.totals.units).toLocaleString();document.querySelector('#cost').textContent=fmt(d.totals.cost_value);document.querySelector('#list').textContent=fmt(d.totals.list_value);document.querySelector('#status').textContent=`Showing ${all.length.toLocaleString()} products${q?' matching search':''}${d.totals.missing_prices?` · ${Number(d.totals.missing_prices).toLocaleString()} items missing a cost or list price`:''}`;render()}
function esc(s){return String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]))}
function render(){const body=document.querySelector('#rows');if(!all.length){body.innerHTML='<tr><td class="empty" colspan="6">No products found.</td></tr>';return}body.innerHTML=all.map(p=>`<tr data-id="${p.id}"><td><div class="part">${esc(p.part_number||'NO PART #')}</div><div class="sub">${esc(p.brand)}</div></td><td><div class="name">${esc(p.name)}</div><div class="sub">${esc(p.category)}</div></td><td>${esc(p.location||'—')}</td><td class="qty"><button data-delta="-1" title="Remove one">−</button><strong class="count">${p.quantity}</strong><button data-delta="1" title="Add one">+</button> <button data-adjust="1" title="Adjust quantity">Edit</button></td><td>${priceInput(p,'purchase_price')}</td><td>${priceInput(p,'list_price')}</td></tr>`).join('');body.querySelectorAll('[data-delta]').forEach(b=>b.onclick=()=>changeQty(b.closest('tr'),Number(b.dataset.delta)));body.querySelectorAll('[data-adjust]').forEach(b=>b.onclick=()=>adjustQty(b.closest('tr')));body.querySelectorAll('[data-price]').forEach(x=>x.onchange=()=>savePrice(x))}
function priceInput(p,k){return `<input class="price" type="number" min="0" step="0.01" value="${p[k]??''}" placeholder="Add" data-price="${k}" aria-label="${k==='purchase_price'?'Purchase':'List'} price for ${esc(p.name)}">`}
async function post(url,body,method='POST'){const r=await fetch(url,{method,headers:{'Content-Type':'application/json','X-CSRF-Token':CSRF},body:JSON.stringify(body)});let d={};try{d=await r.json()}catch{}if(!r.ok)throw Error(d.error||'Save failed');return d}
async function changeQty(row,delta){try{const d=await post(`/api/products/${row.dataset.id}/quantity`,{delta,note:'Quick quantity adjustment'});row.querySelector('.count').textContent=d.quantity;loadTotalsOnly()}catch(e){toast(e.message)}}
async function adjustQty(row){const current=Number(row.querySelector('.count').textContent);const val=prompt(`Set the new quantity (currently ${current}).`,String(current));if(val===null)return;const n=Number(val);if(!Number.isInteger(n)||n<0){toast('Enter a whole number of zero or more.');return}try{if(n===current)return;const d=await post(`/api/products/${row.dataset.id}/quantity`,{delta:n-current,note:'Set quantity'});row.querySelector('.count').textContent=d.quantity;loadTotalsOnly()}catch(e){toast(e.message)}}
async function loadTotalsOnly(){await load()}
async function savePrice(input){try{await post(`/api/products/${input.closest('tr').dataset.id}`,{[input.dataset.price]:input.value},'PATCH');toast('Price saved.');load()}catch(e){toast(e.message)}}
function toast(t){const el=document.querySelector('#toast');el.textContent=t;el.style.display='block';setTimeout(()=>el.style.display='none',2600)}
let timer;document.querySelector('#q').oninput=()=>{clearTimeout(timer);timer=setTimeout(load,220)};document.querySelector('#refresh').onclick=load;document.querySelector('#logout').onclick=e=>{e.currentTarget.form.querySelector('[name=csrf]').value=CSRF};load();
const modal=document.querySelector('#modal'),video=document.querySelector('#video'),canvas=document.querySelector('#canvas'),photo=document.querySelector('#photo');let stream=null;document.querySelector('#scan').onclick=async()=>{modal.classList.add('open');photo.style.display='none';try{stream=await navigator.mediaDevices.getUserMedia({video:{facingMode:{ideal:'environment'}},audio:false});video.srcObject=stream;video.style.display='block';document.querySelector('#capture').style.display='inline-block'}catch(e){video.style.display='none';document.querySelector('#capture').style.display='none';document.querySelector('#scanstatus').textContent='Camera unavailable. Choose an image file instead.'}};function stop(){if(stream)stream.getTracks().forEach(t=>t.stop());stream=null;video.srcObject=null}document.querySelector('#close').onclick=()=>{modal.classList.remove('open');stop()};modal.onclick=e=>{if(e.target===modal){modal.classList.remove('open');stop()}};document.querySelector('#capture').onclick=()=>{canvas.width=video.videoWidth;canvas.height=video.videoHeight;canvas.getContext('2d').drawImage(video,0,0);photo.src=canvas.toDataURL('image/jpeg',.9);photo.style.display='block';stop()};document.querySelector('#file').onchange=e=>{if(e.target.files[0]){photo.src=URL.createObjectURL(e.target.files[0]);photo.dataset.file='1';photo.style.display='block';stop()}};document.querySelector('#ocr').onclick=async()=>{if(!photo.src){document.querySelector('#scanstatus').textContent='Take a photo or choose an image first.';return}if(!window.Tesseract){document.querySelector('#scanstatus').textContent='Text scan library did not load. You can type the part number in search.';return}document.querySelector('#scanstatus').textContent='Reading image…';try{const {data}=await Tesseract.recognize(photo.src,'eng',{logger:m=>{if(m.status==='recognizing text')document.querySelector('#scanstatus').textContent=`Reading image… ${Math.round(m.progress*100)}%`}});let words=data.text.split(/\s+/).map(s=>s.replace(/[^A-Za-z0-9./_-]/g,'')).filter(s=>s.length>=3);words.sort((a,b)=>(/\d/.test(b)-/\d/.test(a))||b.length-a.length);const detected=words[0]||data.text.trim().split(/\s+/)[0]||'';document.querySelector('#q').value=detected;document.querySelector('#scanstatus').textContent=detected?`Detected “${detected}”. Review it in search.`:'No part number found; try a closer, sharper photo.';load()}catch(e){document.querySelector('#scanstatus').textContent='Could not read this image. Try a clearer photo.'}};
</script></body></html>'''

if DB:
    init_db()

if __name__=='__main__':
    app.run(host='0.0.0.0',port=int(os.getenv('PORT','10000')),debug=False)
